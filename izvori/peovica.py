"""Peovica d.o.o. (Omiš): Omiš, Dugi Rat, Šestanovac, Zadvarje, one zone per settlement group.

    python3 -m izvori.peovica [--year 2026]

The home page (or the WordPress media list) links the year workbook "YYYY-Raspored-prikupljanja-otpada….xlsx":
16 sheets, one per group of settlements (full name in row 4), each a 12-month calendar of real date cells
read with openpyxl (cached values, so the week-start array formulas read as dates too). The waste type is
the cell fill: theme 4 tint +0.60 mixed waste, FF00B0F0 paper/cardboard, FFFF00 packaging, FF92D050
"DUBOKA – MKO" (deep containers, one sheet); the legend swatches in rows 5-6 are checked against these
colours. Dates of a neighbouring month inside a month block (greyed out by conditional formatting) are
skipped; every day of the year must appear exactly once in each sheet and sit in its weekday column. The
calendar marks no holiday shifts (Peovica announces them in news posts), so the dates are used as listed.
"""
import argparse
import datetime as dt
import json
import re
import sys
import tempfile
import warnings
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import openpyxl

import podaci
import pravila
from izvori.slavonski_brod_komunalac import fetch

SLUG = "peovica"
SITE = "https://peovica.hr"
MEDIA = SITE + "/wp-json/wp/v2/media?search=raspored&per_page=50&_fields=source_url,date"
MONTHS = ["SIJEČANJ", "VELJAČA", "OŽUJAK", "TRAVANJ", "SVIBANJ", "LIPANJ", "SRPANJ", "KOLOVOZ", "RUJAN",
          "LISTOPAD", "STUDENI", "PROSINAC"]
HEAD = ["po", "ut", "sr", "če", "pe", "su", "ne"]
DAN = ["ponedjeljak", "utorak", "srijeda", "četvrtak", "petak", "subota", "nedjelja"]
# fill -> (code, word the legend label must contain)
FILLS = {("theme", 4, 0.6): ("M", "MIJEŠANI"), ("rgb", "FF00B0F0"): ("K", "PAPIR"),
         ("rgb", "FFFFFF00"): ("P", "AMBALAŽ"), ("rgb", "FF92D050"): ("D", "DUBOKA")}
EMPTY = {None, ("theme", 0, 0.0)}
WORDS = {"PONEDJELJAK": 0, "UTORAK": 1, "SRIJEDA": 2, "SRIJEDU": 2, "ČETVRTAK": 3, "PETAK": 4, "SUBOTA": 5,
         "SUBOTU": 5}
JLS_OF = {
    "Dugi Rat": ["Orij", "Mali Rat", "Sumpetar", "Suhi Potok", "Krilo", "Bajnice", "Duće", "Dugi Rat"],
    "Šestanovac": ["Katuni", "Kreševo", "Šestanovac", "Žeževica", "Grabovac"],
    "Zadvarje": ["Zadvarje"],
    "Omiš": ["Omiš Priko (Vukovarska ulica)", "Lisičina", "Put Vrila", "Zakučac (naselje)", "Borak", "Slavinj",
             "Brzet", "Garma", "Ravnice", "Balića Rat", "Stanići", "Čelina", "Nemira", "Stara Sela Lokva",
             "Lokva Rogoznica", "Medići", "Mimice", "Marušići", "Pisak", "Podašpilje", "Svinišće", "Kučiće",
             "Slime", "Naklice", "Tugare", "Dubrava", "Gata", "Čišla", "Ostrvica", "Zvečanje", "Smolonje",
             "Kostanje", "Podgrađe", "Seoca", "Blato na Cetini", "Trnbusi", "Gornji Dolac", "Srijane", "Putišići",
             "Donji Dolac", "Nova Sela"],
}
PROVIDER = {
    "davatelj": "Peovica d.o.o.",
    "web": SITE,
    "zupanija": "Splitsko-dalmatinska",
    "jls": ["Omiš", "Dugi Rat", "Šestanovac", "Zadvarje"],
    "nazivi": {"P": "Ambalažni otpad (plastična ambalaža)"},
    "napomene": [
        "Uži centar Omiša i Dugog Rata koristi polupodzemne spremnike, pa za njih nema rasporeda odvoza od vrata "
        "do vrata; raspored vrijedi za naselja navedena u zonama.",
        "Kalendar ne označava pomake zbog blagdana; Peovica ih objavljuje u obavijestima na peovica.hr, pa ih "
        "pratite ondje. Datumi su preneseni kako su u kalendaru (npr. 25.12. nema odvoza).",
        "Reciklabilni otpad prikuplja se i prema rasporedu \"Od kvarta do kvarta\" (peovica.hr).",
        "Kontakt: Peovica d.o.o., Vladimira Nazora 12, Omiš, 021 862 388.",
    ],
}


def fill(cell):
    if not cell.fill or not cell.fill.fill_type:
        return None
    c = cell.fill.fgColor
    if c.type == "rgb":
        return ("rgb", c.rgb)
    if c.type == "theme":
        return ("theme", c.theme, round(c.tint, 2))
    return (c.type, c.value)


def text(v):
    return " ".join(str(v).split()) if isinstance(v, str) else ""


def naslov(name):
    """'BLATO NA CETINI' -> 'Blato na Cetini'."""
    return " ".join(w.lower() if w.lower() in ("na", "i") else w.capitalize() for w in name.split())


def settlements(group):
    """Row-4 group name -> settlement names (as in JLS_OF)."""
    if group.upper().startswith("OMIŠ PRIKO (VUKOVARSKA ULICA)"):
        group = group.replace("(Vukovarska ulica)", "(Vukovarska ulica) -", 1)
    out = []
    for part in re.split(r"\s*-\s*", group):
        part = " ".join(part.split())
        if part.upper() == "NASELJE" and out:
            out[-1] += " (naselje)"
        elif part:
            out.append(naslov(part).replace("(vukovarska Ulica)", "(Vukovarska ulica)"))
    return out


def legend(ws):
    """{fill: label} from the swatches in rows 5-6 (label = text after the swatch in row 6 plus text above it)."""
    out = {}
    swatches = [c for c in ws[6] if fill(c) not in EMPTY and not text(c.value)]
    for i, sw in enumerate(swatches):
        end = swatches[i + 1].column if i + 1 < len(swatches) else ws.max_column + 1
        words = [text(c.value) for r in (5, 6) for c in ws[r] if sw.column <= c.column < end and text(c.value)]
        out[fill(sw)] = " ".join(words)
    return out


def read_sheet(ws, year, problems):
    """{date: codes} of one sheet; codes from FILLS ('D' = deep containers)."""
    name = ws.title.strip()
    heads = [(c.row, c.column, MONTHS.index(text(c.value).upper())) for row in ws.iter_rows(max_row=40)
             for c in row if text(c.value).upper() in MONTHS]
    if sorted(m for *_, m in heads) != list(range(12)):
        problems.append(f"{name}: mjeseci {sorted(m + 1 for *_, m in heads)}")
        return {}, {}
    labels = legend(ws)
    found, seen = {}, Counter()
    for r, col, mi in heads:
        if [text(ws.cell(r + 1, col + i).value).lower() for i in range(7)] != HEAD:
            problems.append(f"{name} {MONTHS[mi]}: nema zaglavlja po..ne")
            continue
        for row in range(r + 2, r + 8):
            for i in range(7):
                cell = ws.cell(row, col + i)
                v = cell.value
                if v is None:
                    continue
                if not isinstance(v, dt.datetime):
                    problems.append(f"{name} {cell.coordinate}: nije datum {v!r}")
                    continue
                d = v.date()
                if d.month != mi + 1 or d.year != year:
                    continue  # other month, greyed out in the sheet
                seen[d] += 1
                if d.weekday() != i:
                    problems.append(f"{name} {cell.coordinate}: {d} nije {DAN[i]}")
                f = fill(cell)
                if f in EMPTY:
                    continue
                if f not in FILLS:
                    problems.append(f"{name} {cell.coordinate}: nepoznata boja {f}")
                    continue
                code, word = FILLS[f]
                if word not in labels.get(f, "").upper():
                    problems.append(f"{name}: boja {f} nije u legendi kao {word} ({labels})")
                found[d] = code
    days = {date(year, 1, 1) + dt.timedelta(n) for n in range((date(year + 1, 1, 1) - date(year, 1, 1)).days)}
    if set(seen) != days or max(seen.values()) != 1:
        problems.append(f"{name}: dana u kalendaru {len(seen)}, ponovljenih {sum(1 for n in seen.values() if n > 1)}")
    return found, labels


def days_text(dates):
    wd = sorted({d.weekday() for d in dates})
    return " i ".join(DAN[w] for w in wd)


def find_xlsx(year):
    html = fetch(SITE + "/").decode("utf-8", "replace")
    links = re.findall(rf'href="([^"]+/uploads/[^"]*{year}[^"]*Raspored[^"]*\.xlsx)"', html, re.I)
    if not links:
        media = json.loads(fetch(MEDIA))
        links = [m["source_url"] for m in sorted(media, key=lambda m: m["date"])
                 if m["source_url"].endswith(".xlsx") and str(year) in m["source_url"].rsplit("/", 1)[1]]
    return links[-1] if links else None


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--year", type=int, default=2026)
    year = ap.parse_args(argv).year
    problems = []
    url = find_xlsx(year)
    if not url:
        sys.exit(f"Nema kalendara za {year} na {SITE}")
    print(f"Kalendar: {url}")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "raspored.xlsx"
        fetch(url, path)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # wmf images, header/footer
            wb = openpyxl.load_workbook(path, data_only=True)

    where = {n: j for j, ns in JLS_OF.items() for n in ns}
    zones, totals = {}, Counter()
    for n, ws in enumerate(wb.worksheets, 1):
        if not any(c.value == year for row in ws.iter_rows(max_row=3) for c in row):
            problems.append(f"{ws.title.strip()}: kalendar nije za {year}")
            continue
        group = next((text(c.value) for c in ws[4] if text(c.value)), ws.title.strip())
        found, labels = read_sheet(ws, year, problems)
        names = settlements(group)
        unknown = [x for x in names if x not in where]
        if unknown or not names:
            problems.append(f"{ws.title.strip()}: nepoznata naselja {unknown or group}")
            continue
        by_code = defaultdict(list)
        for d, code in found.items():
            by_code[code].append(d)
        notes = []
        if by_code.get("D"):
            if sorted(by_code["D"]) != pravila.tjedno(year, "sri"):
                problems.append(f"{group}: DUBOKA - MKO nije svaka srijeda")
            notes.append("Duboki (polupodzemni) spremnici za miješani otpad prazne se svake srijede (oznaka "
                         "\"DUBOKA – MKO\" u kalendaru); datumi u rasporedu su za odvoz od vrata do vrata.")
        for code, word in (("K", "papir i karton"), ("P", "ambalažni otpad")):
            if not by_code.get(code):
                label = next((l for f, l in labels.items() if FILLS.get(f, ("",))[0] == code), "")
                notes.append(f"U kalendaru nisu označeni datumi za {word}"
                             + (f" (legenda: \"{label}\")" if label else "") + "; pitajte Peovicu.")
        for f, label in labels.items():  # weekday in the legend text vs the coloured dates
            code = FILLS.get(f, ("",))[0]
            said = {i for w, i in WORDS.items() if re.search(rf"\b{w}\b", label.upper())}
            got = {d.weekday() for d in by_code.get(code, [])}
            if said and got and not got <= said:
                print(f"   UPOZORENJE {group}: legenda \"{label}\", a datumi su {days_text(by_code[code])}")
                notes.append(f"Legenda kalendara navodi \"{label.capitalize()}\", a označeni datumi su "
                             f"{days_text(by_code[code])}; preneseni su označeni datumi.")
        sub = " ".join(labels.get(f, "") for f in labels if FILLS.get(f, ("",))[0] == "K")
        if re.search(r"SLAVINJ\s*-\s*STANIĆI", sub, re.I):
            notes.append("Papir i karton odvozi se samo u dijelu Slavinj – Stanići (prema legendi kalendara).")
        # checks
        per_month = Counter((d.month, c) for d, c in found.items())
        for m in range(1, 13):
            if not 3 <= per_month[(m, "M")] <= 10:  # 3: a holiday without collection
                problems.append(f"{group}: miješani {per_month[(m, 'M')]} puta u {m}. mjesecu")
            if per_month[(m, "K")] > 3 or per_month[(m, "P")] > 5:
                problems.append(f"{group}: {m}. mjesec papir {per_month[(m, 'K')]}, ambalaža {per_month[(m, 'P')]}")
        for code, ds in by_code.items():
            if len({d.weekday() for d in ds}) > 2:
                problems.append(f"{group}: {code} na danima {days_text(ds)}")
            if code in "KP" and len(ds) < 12:
                problems.append(f"{group}: {code} samo {len(ds)} puta")
        rows = [(d, c, False) for d, c in found.items() if c != "D"]
        totals.update(c for _, c, _ in rows)
        parts = defaultdict(list)
        for x in names:
            parts[where[x]].append(x)
        desc = "; ".join(f"{w} {days_text(by_code[c])}" for c, w in (("M", "miješani"), ("K", "papir"),
                                                                     ("P", "ambalaža")) if by_code.get(c))
        for i, (jls, ns) in enumerate(parts.items()):
            key = str(n) + ("AB"[i] if len(parts) > 1 else "")
            zone = {"jls": jls, "podrucje": f"{', '.join(ns)} – {desc}", "opis": group, "ulice": ns}
            other = [f"{', '.join(v)} ({j})" for j, v in parts.items() if j != jls]
            znotes = notes + ([f"Isti raspored vrijedi i za: {'; '.join(other)}."] if other else [])
            if znotes:
                zone["napomena"] = " ".join(znotes)
            zone["raw"] = {str(year): podaci.month_lines(rows)}
            zones[key] = zone

    for p in problems:
        print(f"   PROBLEM {p}")
    if problems or not zones:
        sys.exit("Ništa nije upisano.")
    data = {**PROVIDER, "izvor": url, "zone": zones}
    data["napomene"] = PROVIDER["napomene"] + [f"Izvor: {url} (poveznica na naslovnici {SITE})."]
    out = podaci.PODACI / f"{SLUG}.json"
    if out.exists():  # keep other years of zones that did not change
        old = podaci.load(out)
        for z, zone in zones.items():
            prev = old.get("zone", {}).get(z)
            if prev and prev.get("ulice") == zone["ulice"]:
                zone["raw"] = {**prev["raw"], **zone["raw"]}
    print(f"Zona: {len(zones)}; odvoza {year} (zbroj po zonama): "
          + ", ".join(f"{c} {n}" for c, n in sorted(totals.items())))
    podaci.save(SLUG, data)
    print(f"Upisano: {out.relative_to(podaci.ROOT)}")


if __name__ == "__main__":
    main()
