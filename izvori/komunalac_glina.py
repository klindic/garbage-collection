"""Glina: Komunalac Glina d.o.o. (komunalac-glina.hr), Excel workbook with one sheet per weekday.

    python3 -m izvori.komunalac_glina [--year 2026]

The page "Gospodarenje otpadom" links the year's "obavijest" workbook (old .xls). It is converted to
.xlsx with LibreOffice (soffice --headless; xlrd is not installed) in a temporary folder and read with
openpyxl. The sheet "PLAN ODVOZA OTPADA, YYYY." lists the streets and settlements of each weekday and the
residential buildings with containers ("Kontejneri stambenih zgrada: ..."); the weekday sheets
("ponedjeljak" ... "petak") give the planned dates: mixed waste every week, recyclables (blue paper and
yellow plastic bin, the same day) once a month. A weekday is one zone; buildings with containers that
are emptied on several weekdays get their own zones (mixed waste only, the plan gives no recyclables
for them). Biowaste is not collected ("nije u planu"). Dates falling on a public holiday are moved to
the next working day (Monday to Friday) as the plan says ("NADOKNADA - SLJEDEĆI RADNI DAN ILI POSEBNOM
OBAVIJESTI ODREĐEN RADNI DAN").
"""
import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

import podaci
import pravila
from izvori.slavonski_brod_komunalac import fetch

SLUG = "komunalac-glina"
SITE = "https://komunalac-glina.hr/wp"
PAGE = SITE + "/gospodarenje-otpadom/"
DAYS = ["PONEDJELJAK", "UTORAK", "SRIJEDA", "ČETVRTAK", "PETAK"]
SHEETS = ["ponedjeljak", "utorak", "srijeda", "četvrtak", "petak"]
DAN = ["ponedjeljak", "utorak", "srijeda", "četvrtak", "petak", "subota", "nedjelja"]
CONTAINERS = "Kontejneri stambenih zgrada:"
PROVIDER = {
    "davatelj": "Komunalac Glina d.o.o.",
    "web": SITE,
    "izvor": PAGE,
    "zupanija": "Sisačko-moslavačka",
    "jls": ["Glina"],
    "nazivi": {"P": "Plastika (žuti spremnik)", "K": "Papir (plavi spremnik)"},
    "bioNapomena": "Biootpad se ne odvozi (nije u planu jer nema kompostane i cjenika za uslugu).",
}
NAPOMENE = [
    "Miješani komunalni otpad odvozi se jednom tjedno, a reciklabilni otpad (plavi i žuti spremnik) jednom "
    "mjesečno, okvirno od 07:00 do 15:30.",
    "Reciklažno dvorište: Glinište 3, Glina (ljeti 07:00 – 15:00, zimi 07:30 – 15:30); mobilno reciklažno "
    "dvorište obilazi naselja četiri puta godišnje prema rasporedu u obavijesti.",
    "Glomazni otpad na zahtjev: 099/211-1146 (do 10 sati), info@komunalac-glina.hr.",
]


def cell_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, str) and re.fullmatch(r"\d{2}\.\d{2}\.\d{4}\.?", v.strip()):
        d, m, y = v.strip().rstrip(".").split(".")
        return date(int(y), int(m), int(d))
    return None


def naslov(item):
    """'S. I A. RADIĆA' -> 'S. i A. Radića', 'MAJA (U DIJELU PROMETNICE D6)' -> 'Maja (u dijelu prometnice D6)'."""
    item = item.replace("Ð", "Đ")
    out, depth = [], 0
    words = item.split()
    for i, w in enumerate(words):
        depth += w.count("(") - w.count(")")
        inner = depth > 0 or w.startswith("(") or w.endswith(")") and "(" not in w
        if any(c.isdigit() for c in w) or not w.isupper():
            out.append(w)
        elif w == "I" and 0 < i < len(words) - 1:
            out.append("i")
        elif w in ("I", "II", "III"):
            out.append(w)
        elif inner or w in ("DR.",):
            out.append(w.lower())
        else:
            out.append(re.sub(r"(^|[-.(])(\w)(\w*)", lambda m: m.group(1) + m.group(2) + m.group(3).lower(), w))
    return " ".join(out)


def split_items(text):
    """'J.POGLEDIĆA (KB 2 I 4 ) VRELSKA , ŠAŠ;' -> ['J.Pogledića (kb 2 i 4 )', 'Vrelska', 'Šaš']."""
    text = re.sub(r"\)\s+(?=[A-ZČĆŽŠĐ])", "), ", text)  # a missing comma after a bracket
    text = re.sub(r"(\w)\(", r"\1 (", text)
    return [naslov(" ".join(p.split())) for p in re.split(r"[,;]", text) if p.strip()]


def rows_of(ws):
    return [[v for v in r if v not in (None, "")] for r in ws.iter_rows(values_only=True)]


def read_plan(ws, problems):
    """{weekday: ([streets], [buildings with containers])} from the plan sheet."""
    out, cur = {}, None
    for r in rows_of(ws):
        if not r or not isinstance(r[0], str):
            continue
        head = r[0].strip()
        if head in DAYS:
            cur = DAYS.index(head)
            out[cur] = ([], [])
            r = r[1:]
        if cur is None or not r:
            continue
        text = " ".join(str(x) for x in r)
        if text.startswith(CONTAINERS):
            out[cur][1].extend(split_items(text[len(CONTAINERS):]))
        elif text != "PLAN ODVOZA OTPADA":
            out[cur][0].extend(s for s in split_items(text) if s not in out[cur][0])
    if sorted(out) != list(range(5)):
        problems.append(f"plan: dani {sorted(out)}")
    return out


def read_day(ws, wd, year, problems):
    """(mixed-waste dates, recyclables dates, bulky waste months) of a weekday sheet."""
    kind, mko, rec, bulky = None, [], [], None
    rows = rows_of(ws)
    for i, r in enumerate(rows):
        first = str(r[0]) if r else ""
        if first.startswith("MIJEŠANI KOMUNALNI OTPAD"):
            kind = mko
        elif first.startswith("RECIKLABILNI KOMUNALNI OTPAD"):
            kind = rec
        elif first.startswith("PLAN PREUZIMANJA GLOMAZNOG") and i + 1 < len(rows):
            bulky = " ".join(str(x) for x in rows[i + 1])
            kind = None
        elif r and not cell_date(r[0]):
            kind = None
        if kind is not None:
            kind += [d for d in map(cell_date, r) if d]
    for name, dates in (("miješani", mko), ("reciklabilni", rec)):
        if any(d.year != year for d in dates):
            problems.append(f"{SHEETS[wd]}: {name} – datumi izvan {year}.")
        if any(d.weekday() != wd for d in dates):
            problems.append(f"{SHEETS[wd]}: {name} – datum nije {DAN[wd]}")
        if len(set(dates)) != len(dates):
            problems.append(f"{SHEETS[wd]}: {name} – datum dvaput")
    if not 51 <= len(mko) <= 53 or len(rec) != 12 or sorted({d.month for d in rec}) != list(range(1, 13)):
        problems.append(f"{SHEETS[wd]}: {len(mko)} odvoza miješanog, {len(rec)} reciklabilnog otpada")
    return mko, rec, bulky


def next_working(d, hol):
    """The plan's rule: a collection on a holiday is made up on the next working day (Monday to Friday)."""
    if d not in hol:
        return d, False
    while d in hol or d.weekday() >= 5:
        d += timedelta(days=1)
    return d, True


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--year", type=int, default=2026)
    year = ap.parse_args(argv).year
    problems = []
    html = fetch(PAGE).decode("utf-8", "replace")
    links = [u for u in re.findall(r'href="([^"]+\.xlsx?)"', html) if str(year) in u.rsplit("/", 1)[1]]
    if not links:
        sys.exit(f"Na {PAGE} nema obavijesti (xls) za {year}.")
    url = links[-1]
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("Treba LibreOffice (soffice) za pretvaranje .xls u .xlsx.")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "obavijest.xls"
        fetch(url, src)
        subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/lo",
                        "--convert-to", "xlsx", "--outdir", tmp, str(src)], check=True, capture_output=True, timeout=180)
        wb = openpyxl.load_workbook(Path(tmp) / "obavijest.xlsx", data_only=True)
    plan_ws = next((ws for ws in wb.worksheets if ws.title.startswith("PLAN ODVOZA OTPADA") and str(year) in ws.title), None)
    if plan_ws is None:
        sys.exit(f"U radnoj knjizi nema lista 'PLAN ODVOZA OTPADA, {year}.': {wb.sheetnames}")
    plan = read_plan(plan_ws, problems)
    days = {}
    for wd, name in enumerate(SHEETS):
        if name not in wb.sheetnames:
            problems.append(f"nema lista '{name}'")
            continue
        days[wd] = read_day(wb[name], wd, year, problems)
        sheet_text = " ".join(str(x) for r in rows_of(wb[name])[:12] for x in r).replace("Ð", "Đ")
        for s in plan.get(wd, ([], []))[0][:3]:
            if s.split()[0].upper().replace("Ð", "Đ") not in sheet_text.upper():
                problems.append(f"list '{name}': nema ulice {s} iz plana")
    hol = set(pravila.blagdani(year))

    zones, moved = {}, {}

    def rows_for(dates, code):
        out = []
        for d in dates:
            new, m = next_working(d, hol)
            if m:
                moved[d] = new
            out.append((new, m, code))
        return out


    def add(jls_desc, ulice, rows, napomena):
        merged = {}
        for d, m, code in rows:
            c, mm = merged.get(d, ("", False))
            if set(code) & set(c):
                problems.append(f"{jls_desc}: {d} {code} dvaput")
            merged[d] = (c + code, mm or m)
        zone = {"jls": "Glina", "podrucje": jls_desc, "ulice": ulice, "napomena": napomena,
                "raw": {str(year): podaci.month_lines([(d, c, m) for d, (c, m) in merged.items()])}}
        zones[str(len(zones) + 1)] = zone
        print(f"Zona {len(zones)}: {len(ulice)} ulica/naselja – {jls_desc}")

    for wd, (streets, _) in sorted(plan.items()):
        if wd not in days:
            continue
        mko, rec, bulky = days[wd]
        add(f"{DAN[wd].capitalize()} – " + ", ".join(streets[:4]) + " …", streets,
            rows_for(mko, "M") + rows_for(rec, "PK"),
            f"Glomazni otpad na zahtjev: {bulky.lower()}." if bulky else None)
    # buildings with containers: emptied on every weekday that lists them
    buildings = defaultdict(list)
    for wd, (_, cont) in plan.items():
        for b in cont:
            buildings[b].append(wd)
    groups = defaultdict(list)
    for b, wds in buildings.items():
        groups[tuple(sorted(wds))].append(b)
    for wds, names in sorted(groups.items()):
        rows = [r for wd in wds if wd in days for r in rows_for(days[wd][0], "M")]
        add("Kontejneri stambenih zgrada – " + " i ".join(DAN[w] for w in wds) + ": " + ", ".join(names),
            [f"{n} (kontejneri stambenih zgrada)" for n in names], rows,
            "Miješani otpad iz kontejnera stambenih zgrada; plan ne navodi odvoz reciklabilnog otpada za kontejnere.")
    for z in zones.values():
        if z.get("napomena") is None:
            del z["napomena"]

    for p in problems:
        print(f"   PROBLEM {p}")
    if problems or not zones:
        sys.exit("Ništa nije upisano.")
    shifts = ", ".join(f"{d:%d.%m.} → {n:%d.%m.}" for d, n in sorted(moved.items()))
    data = {**PROVIDER, "napomene": NAPOMENE + [
        "Blagdani: prema planu se odvoz s blagdana nadoknađuje sljedeći radni dan ili drugog dana prema posebnoj "
        f"obavijesti; uzet je sljedeći radni dan (pon – pet): {shifts}",
        f"Izvor: {url}",
    ], "zone": zones}
    path = podaci.PODACI / f"{SLUG}.json"
    if path.exists():  # keep other years of zones that did not change
        old = podaci.load(path)
        for z, zone in zones.items():
            prev = old.get("zone", {}).get(z)
            if prev and prev.get("ulice") == zone["ulice"]:
                zone["raw"] = {**prev["raw"], **zone["raw"]}
    podaci.save(SLUG, data)
    print(f"Upisano: {path.relative_to(podaci.ROOT)} ({len(zones)} zona)")


if __name__ == "__main__":
    main()
