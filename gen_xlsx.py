#!/usr/bin/env python3
"""Write one Excel workbook per zone and year (excel/Raspored_odvoza_ZonaN_<year>.xlsx) from the schedule data in odvoz.html.

Needs openpyxl. Run it after changing the schedule; build.sh only copies the committed files.
"""
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gen_ics import ORDER, load_data, parse_schedule

TYPES = {  # name, fill, font colour, column width in "Raspored"
    "M": ("Miješani komunalni otpad", "262626", "FFFFFF", 16),
    "B": ("Biootpad", "7B4A2D", "FFFFFF", 12),
    "P": ("Plastika, staklo i metal", "FFE600", "000000", 16),
    "K": ("Papir i karton", "3B6EF5", "FFFFFF", 13),
}
MONTHS = ["siječanj", "veljača", "ožujak", "travanj", "svibanj", "lipanj", "srpanj",
          "kolovoz", "rujan", "listopad", "studeni", "prosinac"]
DAYS = ["ponedjeljak", "utorak", "srijeda", "četvrtak", "petak", "subota", "nedjelja"]
DAYS_INS = ["ponedjeljkom", "utorkom", "srijedom", "četvrtkom", "petkom", "subotom", "nedjeljom"]
MOVED = "Pomaknuto zbog neradnog dana u tjednu"
PILOT = "Testni projekt: plastika 2x mjesečno"

HEAD_FILL = "375623"
EDGE = Side(style="thin", color="BFBFBF")
BORDER = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATE_FMT = "DD.MM.YYYY."


def put(ws, ref, value, bold=False, size=None, color=None, fill=None, center=False, border=True, fmt=None, **font):
    c = ws[ref]
    c.value = value
    c.font = Font(name="Arial", bold=bold, sz=size, color=color, **font)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if center:
        c.alignment = CENTER
    if border:
        c.border = BORDER
    if fmt:
        c.number_format = fmt
    return c


def head(ws, ref, value, fill=HEAD_FILL, color="FFFFFF"):
    put(ws, ref, value, bold=True, color=color, fill=fill, center=True)


def notes(data, zone, rows, kinds):
    z = data["zones"][zone]
    regular = DAYS_INS[DAYS.index(regular_day(rows))]
    out = [f'Crveno označeni datumi u listu "Raspored" su pomaknuti zbog neradnog dana u tom tjednu (inače je odvoz {regular}).']
    if z.get("note"):
        out.append(z["note"])
    if "B" not in kinds:
        out.append("Biootpad se na ovom području ne odvozi; preporučuje se kućno kompostiranje.")
    else:
        out.append("Biootpad samo za korisnike koji su odabrali predaju biootpada u spremnicima."
                   + (f" Ne odvozi se u {z['noBioIn']}." if z.get("noBioIn") else ""))
    out.append("Spremnike iznijeti na javnu površinu najkasnije do 07:00.")
    out.append('Reciklažna dvorišta: "Sisak Stari" (Kralja Zvonimira 7B) i "Novi Sisak" (Capraška 4), pon-pet 08-20, sub 08-13.')
    return out


def regular_day(rows):
    days = [DAYS[d.weekday()] for d, _, moved in rows if not moved]
    return max(DAYS, key=days.count)


def workbook(data, zone, year):
    z = data["zones"][zone]
    rows = [r for r in parse_schedule(data, zone) if r[0].year == year]
    kinds = [t for t in ORDER if any(t in types for _, types, _ in rows)]
    pilot_from = z.get("pilotFrom")
    last = 4 + len(rows)
    tcol = {t: get_column_letter(4 + i) for i, t in enumerate(kinds)}  # type columns in "Raspored"
    note_col = get_column_letter(4 + len(kinds))

    wb = Workbook()
    ov = wb.active
    ov.title = "Pregled"
    sch = wb.create_sheet("Raspored")
    by = wb.create_sheet("Po vrsti")

    # Raspored: one row per collection day, "x" under each bin type.
    put(sch, "A1", f"Raspored odvoza otpada {year}., Zona {zone} ({z['place']})", bold=True, size=14, border=False)
    put(sch, "A2", '"x" = odvoz tog dana. Spremnike iznijeti do 07:00. '
                   "Prošli datumi su posivljeni.", size=9, color="595959", italic=True, border=False)
    for col, title in zip("ABC", ("Datum", "Dan", "Mjesec")):
        head(sch, f"{col}4", title)
    for t in kinds:
        head(sch, f"{tcol[t]}4", TYPES[t][0], fill=TYPES[t][1], color=TYPES[t][2])
    head(sch, f"{note_col}4", "Napomena")
    sch.row_dimensions[4].height = 32
    for r, (day, types, moved) in enumerate(rows, start=5):
        red = {"bold": True, "color": "C00000"} if moved else {}
        put(sch, f"A{r}", datetime(day.year, day.month, day.day), center=True, fmt=DATE_FMT, **red)
        put(sch, f"B{r}", DAYS[day.weekday()], center=True, **red)
        put(sch, f"C{r}", MONTHS[day.month - 1], center=True)
        for t in kinds:
            put(sch, f"{tcol[t]}{r}", "x" if t in types else None, center=True)
        note = [MOVED] if moved else []
        if pilot_from and "P" in types and day.isoformat() >= pilot_from:
            note.append(PILOT)
        put(sch, f"{note_col}{r}", "; ".join(note) or None).alignment = Alignment(vertical="center")
    for t in kinds:
        rng = f"{tcol[t]}5:{tcol[t]}{last}"
        fill = PatternFill("solid", fgColor=TYPES[t][1], bgColor=TYPES[t][1])
        sch.conditional_formatting.add(rng, FormulaRule(formula=[f'{tcol[t]}5="x"'],
                                                        font=Font(bold=True, color=TYPES[t][2]), fill=fill))
    sch.conditional_formatting.add(f"A5:{note_col}{last}",
                                   FormulaRule(formula=["$A5<TODAY()"], font=Font(color="A6A6A6"), stopIfTrue=False))
    sch.freeze_panes = "A5"
    sch.auto_filter.ref = f"A4:{note_col}{last}"
    for col, width in zip("ABC", (13, 12, 11)):
        sch.column_dimensions[col].width = width
    for t in kinds:
        sch.column_dimensions[tcol[t]].width = TYPES[t][3]
    sch.column_dimensions[note_col].width = 44

    # Po vrsti: one row per bin per day, handy for filtering by type.
    for col, title in zip("ABCD", ("Datum", "Dan", "Vrsta otpada", "Napomena")):
        head(by, f"{col}1", title)
    r = 1
    for day, types, moved in rows:
        for t in types:
            r += 1
            put(by, f"A{r}", datetime(day.year, day.month, day.day), fmt=DATE_FMT)
            put(by, f"B{r}", DAYS[day.weekday()])
            put(by, f"C{r}", TYPES[t][0], bold=True, color=TYPES[t][2], fill=TYPES[t][1])
            put(by, f"D{r}", MOVED if moved else None)
    by.freeze_panes = "A2"
    by.auto_filter.ref = f"A1:D{r}"
    for col, width in zip("ABCD", (13, 12, 26, 40)):
        by.column_dimensions[col].width = width

    # Pregled: next date per type relative to TODAY(), counts per month.
    rs = lambda col: f"Raspored!${col}$5:${col}${last}"  # noqa: E731
    put(ov, "A1", f"Zona {zone}, {year}: pregled odvoza", bold=True, size=14, border=False)
    put(ov, "A3", "Danas", bold=True, border=False)
    put(ov, "B3", "=TODAY()", border=False, fmt=DATE_FMT)
    for col, title in zip("ABCDE", ("Vrsta otpada", "Sljedeći odvoz", "Dan", "Za dana", f"Ukupno u {year}.")):
        head(ov, f"{col}5", title)
    weekdays = ",".join(f'"{d}"' for d in ["nedjelja"] + DAYS[:6])
    for r, t in enumerate(kinds, start=6):
        put(ov, f"A{r}", TYPES[t][0], bold=True, color=TYPES[t][2], fill=TYPES[t][1])
        put(ov, f"B{r}", f'=IFERROR(1/(1/_xlfn.MINIFS({rs("A")},{rs(tcol[t])},"x",{rs("A")},">="&$B$3)),'
                         f'"nema više u {year}.")', center=True, fmt=DATE_FMT)
        put(ov, f"C{r}", f'=IF(ISNUMBER(B{r}),INDEX({{{weekdays}}},WEEKDAY(B{r})),"")', center=True)
        put(ov, f"D{r}", f'=IF(ISNUMBER(B{r}),B{r}-$B$3,"")', center=True, fmt="0")
        put(ov, f"E{r}", f'=COUNTIF({rs(tcol[t])},"x")', center=True)
    top = 6 + len(kinds) + 1
    put(ov, f"A{top}", "Broj odvoza po mjesecu", bold=True, size=12, border=False)
    head(ov, f"A{top + 1}", "Mjesec")
    ov.row_dimensions[top + 1].height = 30
    for i, t in enumerate(kinds):
        head(ov, f"{get_column_letter(2 + i)}{top + 1}", TYPES[t][0], fill=TYPES[t][1], color=TYPES[t][2])
    for m, month in enumerate(MONTHS):
        r = top + 2 + m
        put(ov, f"A{r}", month)
        for i, t in enumerate(kinds):
            put(ov, f"{get_column_letter(2 + i)}{r}",
                f'=COUNTIFS({rs("C")},$A{r},{rs(tcol[t])},"x")', center=True)
    total = top + 2 + len(MONTHS)
    put(ov, f"A{total}", "Ukupno", bold=True)
    for i, _ in enumerate(kinds):
        col = get_column_letter(2 + i)
        put(ov, f"{col}{total}", f"=SUM({col}{top + 2}:{col}{total - 1})", bold=True, center=True)
    put(ov, f"A{total + 2}", "Napomene:", bold=True, size=10, border=False)
    for i, text in enumerate(notes(data, zone, rows, kinds)):
        put(ov, f"A{total + 3 + i}", f"• {text}", size=10, border=False)
    for col, width in zip("ABCDE", (26, 24, 24, 24, 20)):
        ov.column_dimensions[col].width = width
    return wb


if __name__ == "__main__":
    here = Path(__file__).parent
    data = load_data((here / "odvoz.html").read_text(encoding="utf-8"))
    out_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else here / "excel"
    out_dir.mkdir(exist_ok=True)
    for zone, z in data["zones"].items():
        for year in z["raw"]:
            workbook(data, zone, int(year)).save(out_dir / f"Raspored_odvoza_Zona{zone}_{year}.xlsx")
